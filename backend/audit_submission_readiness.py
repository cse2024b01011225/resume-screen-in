import sys
from pathlib import Path
import io
import time

backend_dir = Path(__file__).resolve().parent
sys.path.insert(0, str(backend_dir))

from fastapi.testclient import TestClient
from app.main import app
import docx
from reportlab.pdfgen import canvas

def build_sample_pdf(name: str, email: str, phone: str, skills_str: str, projects_str: str, edu_str: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 750, name)
    c.drawString(72, 730, f"Email: {email} | Phone: {phone} | github.com/{name.lower().replace(' ', '')}")
    c.drawString(72, 700, f"Education: {edu_str}")
    c.drawString(72, 670, f"Technical Skills: {skills_str}")
    c.drawString(72, 640, f"Academic Projects: {projects_str}")
    c.drawString(72, 610, "Internship Experience: Web Development Intern at Startup Labs.")
    c.drawString(72, 580, "Certifications: Certified Python Specialist via Coursera.")
    c.save()
    return buf.getvalue()

def build_sample_docx(name: str, email: str, phone: str, skills_str: str, projects_str: str, edu_str: str) -> bytes:
    doc = docx.Document()
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"Email: {email} | Phone: {phone} | linkedin.com/in/{name.lower().replace(' ', '')}")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(edu_str)
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(skills_str)
    doc.add_heading("Projects & Experience", level=1)
    doc.add_paragraph(projects_str)
    doc.add_paragraph("Internships: Summer Intern at GlobalSoft Inc.")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def run_comprehensive_audit():
    print("=" * 70)
    print("FINAL SUBMISSION READINESS AUDIT: RESUME SCREEN-IN")
    print("=" * 70)

    with TestClient(app) as client:
        # TEST 1: Root & Health Check + Developer Attribution
        print("\n[CHECK 1] Root Discovery & Developer Attribution...")
        res = client.get("/")
        assert res.status_code == 200, f"Root endpoint failed: {res.text}"
        root_data = res.json()
        assert root_data.get("service") == "Resume Screen-In", f"Unexpected service: {root_data.get('service')}"
        assert root_data.get("developer") == "Yash Tomar", f"Unexpected developer: {root_data.get('developer')}"
        assert root_data.get("category") == "HCL Capstone Project"
        print(f"  -> PASS: Service: {root_data['service']} | Developer: {root_data['developer']} | Category: {root_data['category']}")

        print("\n[CHECK 2] Backend Health & Database Connectivity...")
        res = client.get("/api/health")
        assert res.status_code == 200
        health_data = res.json()
        assert health_data["status"] == "healthy"
        assert health_data["database"] == "connected"
        print(f"  -> PASS: Backend status: {health_data['status']} | SQLite: {health_data['database']}")

        # TEST 3: Job Creation
        print("\n[CHECK 3] Job Requisition Creation...")
        timestamp = int(time.time() * 1000)
        job_payload = {
            "title": f"Full-Stack Python Engineer #{timestamp}",
            "description": "Looking for a Python backend engineer proficient in FastAPI, React, SQLite, Git, and Docker. Experience with REST APIs and capstone projects desired.",
            "required_skills": ["Python", "FastAPI", "React", "SQLite", "Git", "Docker"],
            "experience_level": "Entry / Fresher",
            "is_fresher_friendly": True,
            "min_education": "Bachelor's Degree",
        }
        res = client.post("/api/jobs", json=job_payload)
        assert res.status_code == 201, f"Failed to create job: {res.text}"
        job_data = res.json()
        job_id = job_data["id"]
        print(f"  -> PASS: Created Job ID #{job_id} ('{job_data['title']}')")

        # TEST 4: Duplicate Requisition Prevention
        print("\n[CHECK 4] Accidental Duplicate Requisition Prevention...")
        dup_res = client.post("/api/jobs", json=job_payload)
        assert dup_res.status_code == 409, f"Duplicate should return 409, got {dup_res.status_code}"
        print(f"  -> PASS: Duplicate submission blocked with 409: {dup_res.json()['detail']}")

        # TEST 5: Batch Multi-Format Upload (PDF + DOCX)
        print("\n[CHECK 5] Batch Multi-Format Resume Ingestion (PDF + DOCX)...")
        # Candidate 1: Yash Tomar (PDF - strong match)
        pdf_bytes = build_sample_pdf(
            name="Yash Tomar",
            email="yash.tomar@example.com",
            phone="+91 9876543210",
            skills_str="Python, FastAPI, React, SQLite, Git, Docker, REST APIs, HTML/CSS.",
            projects_str="Resume Screen-In: Built complete full-stack web platform with React, Vite, and FastAPI.",
            edu_str="B.Tech Computer Science and Engineering, CGPA 8.7/10."
        )

        # Candidate 2: Ritu Sen (DOCX - weak match, different tech stack)
        docx_bytes = build_sample_docx(
            name="Ritu Sen",
            email="ritu.sen@example.com",
            phone="+91 9123456780",
            skills_str="Java, Spring Boot, Hibernate, MySQL, Angular.",
            projects_str="Enterprise Banking Portal: Developed banking microservices in Java Spring Boot.",
            edu_str="Bachelor of Science in Information Technology."
        )

        files = [
            ("files", ("yash_tomar_resume.pdf", pdf_bytes, "application/pdf")),
            ("files", ("ritu_sen_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ]
        upload_res = client.post(f"/api/jobs/{job_id}/resumes", files=files)
        assert upload_res.status_code == 200, f"Upload failed: {upload_res.text}"
        screened = upload_res.json()
        assert len(screened) == 2, f"Expected 2 candidates, got {len(screened)}"
        print(f"  -> PASS: Successfully parsed and screened {len(screened)} resumes (1 PDF, 1 DOCX).")

        # TEST 6: Real Dynamic Matching & Ranking Verification
        print("\n[CHECK 6] Evidence-Based Dynamic Scoring & Candidate Ranking...")
        cand_yash = next(c for c in screened if "Yash" in c["candidate_name"])
        cand_ritu = next(c for c in screened if "Ritu" in c["candidate_name"])

        print(f"  Candidate 1 (Yash Tomar - Strong Match): Score = {cand_yash['match_score']}%")
        print(f"    Matching Skills: {cand_yash['matching_skills']}")
        print(f"    Missing Evidence: {cand_yash['missing_evidence']}")
        print(f"    Evidence Quotes Count: {len(cand_yash['evidence_snippets'])}")
        print(f"    Fresher Highlights Count: {len(cand_yash['fresher_highlights'])}")

        print(f"  Candidate 2 (Ritu Sen - Weak Match): Score = {cand_ritu['match_score']}%")
        print(f"    Matching Skills: {cand_ritu['matching_skills']}")
        print(f"    Missing Evidence: {cand_ritu['missing_evidence']}")

        # Assert dynamic, real job-specific match scoring
        assert cand_yash["match_score"] > cand_ritu["match_score"], "Matching candidate score should be strictly greater than mismatched candidate"
        assert cand_yash["match_score"] >= 80.0, "Strong match candidate should score >= 80%"
        assert cand_ritu["match_score"] < 65.0, "Mismatched candidate should score lower"
        assert len(cand_yash["evidence_snippets"]) > 0, "Candidate must have real extracted evidence snippets"
        assert len(cand_yash["missing_evidence"]) == 0, "Yash had all required skills"
        assert len(cand_ritu["missing_evidence"]) > 0, "Ritu should have missing evidence for FastAPI/React/Docker"
        print("  -> PASS: Scores are dynamically calculated from real text extraction.")

        # TEST 7: Candidate Analysis Deep-Dive
        print("\n[CHECK 7] Candidate Analysis Endpoint Verification...")
        detail_res = client.get(f"/api/candidates/{cand_yash['id']}")
        assert detail_res.status_code == 200
        detail = detail_res.json()
        assert detail["candidate_name"] == "Yash Tomar"
        assert detail["candidate_email"] == "yash.tomar@example.com"
        assert "score_breakdown" in detail
        assert detail["score_breakdown"]["technical_score"] > 0
        assert detail["score_breakdown"]["experience_project_score"] > 0
        assert detail["score_breakdown"]["education_score"] > 0
        print(f"  -> PASS: Detailed analysis retrieved with complete 3-pillar breakdown.")

        # TEST 8: Candidate Shortlisting Toggle
        print("\n[CHECK 8] Shortlist Status Persistence...")
        shortlist_res = client.patch(f"/api/candidates/{cand_yash['id']}/shortlist", json={"is_shortlisted": True})
        assert shortlist_res.status_code == 200
        assert shortlist_res.json()["is_shortlisted"] is True
        # Re-fetch from DB to verify persistence
        verify_res = client.get(f"/api/candidates/{cand_yash['id']}")
        assert verify_res.json()["is_shortlisted"] is True
        print(f"  -> PASS: Candidate successfully shortlisted and verified in database.")

        # TEST 9: CSV Export Verification
        print("\n[CHECK 9] CSV Export Generation...")
        csv_res = client.get(f"/api/jobs/{job_id}/export-csv")
        assert csv_res.status_code == 200
        assert "text/csv" in csv_res.headers["content-type"]
        csv_content = csv_res.text
        assert "Yash Tomar" in csv_content
        assert "Ritu Sen" in csv_content
        assert "YES - SHORTLISTED" in csv_content
        print(f"  -> PASS: CSV export generated with valid headers and candidate records.")

        # TEST 10: Safe Job Deletion & Candidate Integrity
        print("\n[CHECK 10] Safe Job Deletion & Cascade Isolation...")
        # Attempt delete without force -> must be blocked
        blocked_del = client.delete(f"/api/jobs/{job_id}")
        assert blocked_del.status_code == 400
        print(f"  -> PASS: Deletion blocked to protect {len(screened)} active candidate records.")

        # Delete with force=true -> must succeed and remove ONLY this job and its candidates
        forced_del = client.delete(f"/api/jobs/{job_id}?force=true")
        assert forced_del.status_code == 200
        assert forced_del.json()["deleted_candidates_count"] == 2

        # Verify job is gone
        assert client.get(f"/api/jobs/{job_id}").status_code == 404
        # Verify candidates are gone
        assert client.get(f"/api/candidates/{cand_yash['id']}").status_code == 404
        assert client.get(f"/api/candidates/{cand_ritu['id']}").status_code == 404
        print(f"  -> PASS: Job #{job_id} and its associated candidates safely removed without collateral damage.")

    print("\n" + "=" * 70)
    print("ALL 10 AUDIT CHECKS PASSED WITH ZERO ERRORS!")
    print("=" * 70)

if __name__ == "__main__":
    run_comprehensive_audit()
